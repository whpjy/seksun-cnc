from __future__ import annotations

import argparse
import hashlib
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROCESS_SHEETS = ("切削加工", "特种加工", "热处理与表面处理", "检验质控")
FAMILY_KEYS = {
    "切削加工": "cutting",
    "特种加工": "special",
    "热处理与表面处理": "heat_surface",
    "检验质控": "quality",
}

# A manufacturing route process can cover several executable CAM strategies.  A
# mapping means "related capability", not that every variant is production-ready.
CAM_MAPPINGS = {
    "GX-C-07": [
        "face_milling",
        "profile_roughing",
        "internal_profile_roughing",
        "pocket_roughing",
        "slot_roughing",
        "adaptive_clearing",
    ],
    "GX-C-08": [
        "profile_contouring",
        "profile_finishing",
        "internal_profile_finishing",
        "pocket_finishing",
        "slot_finishing",
    ],
    "GX-C-11": ["drilling"],
    "GX-C-16": ["helical_boring"],
    "GX-C-09": ["surface_roughing", "surface_3d", "waterline"],
    "GX-C-41": ["edge_chamfer", "tab_removal"],
}


def _clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def _records(sheet: Any, header_row: int = 1) -> list[dict[str, Any]]:
    headers = [_clean(cell.value) for cell in sheet[header_row]]
    result: list[dict[str, Any]] = []
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        record = {
            str(header): _clean(value)
            for header, value in zip(headers, row, strict=False)
            if header is not None and _clean(value) is not None
        }
        if record:
            result.append(record)
    return result


def _value(record: dict[str, Any], *candidates: str) -> Any:
    for candidate in candidates:
        if candidate in record:
            return record[candidate]
    return None


def _process_record(record: dict[str, Any], sheet_name: str, application: dict[str, Any]) -> dict[str, Any]:
    code = str(_value(record, "工序编码", "工序代码") or "").strip()
    joined_text = " ".join(str(value) for value in (*record.values(), *application.values()))
    evidence_level = "trial_required" if any(token in joined_text for token in ("待验证", "经验", "试切")) else "reference"
    cam_ids = CAM_MAPPINGS.get(code, [])
    return {
        "code": code,
        "family": FAMILY_KEYS[sheet_name],
        "family_name": sheet_name,
        "category": _value(record, "工序大类", "类别", "工序类别"),
        "name": _value(record, "工序名称", "名称"),
        "description": _value(record, "工序定义与说明", "工艺说明", "工序说明", "说明"),
        "equipment": _value(record, "典型设备", "设备"),
        "capability": _value(record, "参考能力/验收指标", "加工能力/适用范围", "加工能力", "适用范围"),
        "roughness_ra": _value(record, "Ra参考值(μm)", "典型粗糙度Ra(μm)", "表面粗糙度Ra(μm)", "粗糙度"),
        "materials": _value(record, "适用材料", "材料"),
        "key_parameters": _value(record, "关键工艺参数", "关键参数", "工艺参数"),
        "quality_controls": _value(record, "质量控制要点", "质量控制"),
        "common_defects": _value(record, "常见缺陷", "缺陷"),
        "inspection": _value(record, "检验方式", "检验方法/量具", "检验方法", "量具"),
        "standards": _value(record, "标准及适用依据", "参考标准", "引用标准"),
        "applicability": {
            "applicable_conditions": _value(application, "适用条件"),
            "prerequisites": _value(application, "前序准备", "前置条件/输入", "前置条件"),
            "tooling_consumables": _value(application, "刀具耗材", "工装/刀具/耗材"),
            "parameter_method": _value(application, "参数确定方法", "关键参数与确定方法", "关键参数"),
            "successors": _value(application, "后序衔接", "后续工序/衔接"),
            "release_requirements": _value(application, "检验与放行", "检验与放行要求"),
            "risk_response": _value(application, "风险处置", "风险与异常处置"),
            "evidence_boundaries": _value(application, "证据与边界", "证据来源与适用边界"),
            "record_type": _value(application, "记录类型", "记录/表单建议"),
        },
        "cam_operation_ids": cam_ids,
        "execution_state": "partially_executable" if cam_ids else "reference_only",
        "evidence_level": evidence_level,
        "approval_policy": "engineer_required",
        "source": {"sheet": sheet_name, "record_code": code},
    }


def build_library(source: Path) -> dict[str, Any]:
    workbook = load_workbook(source, data_only=True, read_only=True)
    application_rows = _records(workbook["工序应用要点"], header_row=2)
    applications = {
        str(_value(row, "工序编码", "工序代码")): row
        for row in application_rows
        if _value(row, "工序编码", "工序代码")
    }
    processes: list[dict[str, Any]] = []
    for sheet_name in PROCESS_SHEETS:
        for record in _records(workbook[sheet_name]):
            code = str(_value(record, "工序编码", "工序代码") or "")
            if code.startswith("GX-"):
                processes.append(_process_record(record, sheet_name, applications.get(code, {})))

    processes.sort(key=lambda item: item["code"])
    reverse_mapping = {
        cam_id: process["code"]
        for process in processes
        for cam_id in process["cam_operation_ids"]
    }
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    family_counts = {
        family: sum(process["family"] == family for process in processes)
        for family in FAMILY_KEYS.values()
    }
    return {
        "schema_version": "1.0.0",
        "catalog_version": "2026.09.13",
        "source": {"filename": source.name, "sha256": digest},
        "summary": {
            "process_count": len(processes),
            "family_counts": family_counts,
            "cam_mapped_process_count": sum(bool(item["cam_operation_ids"]) for item in processes),
            "application_record_count": len(applications),
        },
        "processes": processes,
        "cam_operation_mapping": reverse_mapping,
        "typical_routes": _records(workbook["典型工艺路线"], header_row=2),
        "tolerance_roughness_rules": _records(workbook["公差粗糙度对照"], header_row=2),
        "allowance_rules": _records(workbook["加工余量参考"], header_row=2),
        "standards_index": _records(workbook["引用标准索引"], header_row=2),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Import the machining process workbook into the runtime knowledge catalog.")
    parser.add_argument("source", type=Path)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "data" / "manufacturing_process_library.json",
    )
    args = parser.parse_args()
    payload = build_library(args.source.resolve())
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Imported {payload['summary']['process_count']} processes to {args.output}")


if __name__ == "__main__":
    main()
